from __future__ import annotations

import csv
import json
import time

import pytest

from utils import scraper
from utils.scraper import ScrapeOptions, export_csv, export_json, export_xlsx, scrape_from_html, scrape_with_selector


def _opts(**kw) -> ScrapeOptions:
    kw.setdefault("robots_check", False)
    return ScrapeOptions(**kw)


def test_extract_value_text_and_attribute():
    el = scraper._soup('<div class="p"><strong>X</strong><a href="/u">L</a></div>').select_one(".p")
    assert scraper._extract_value(el, "strong", "text", strip_html=True) == "X"
    assert scraper._extract_value(el, "a@href", "text", strip_html=True) == "/u"
    assert scraper._extract_value(el, "yok", "text", strip_html=True) == ""


def test_table_mode_with_column_mapping(site):
    cols = {"name": "strong", "price": ".price", "stock": ".stock", "link": "a@href"}
    r = scrape_with_selector(f"{site.base}/urunler.html", ".product", "table", options=_opts(columns=cols))
    assert r.title == "Örnek Mağaza — Ürünler"
    assert r.columns == ["Ürün / başlık", "Fiyat / değer", "Stok / meta", "Bağlantı"]
    assert r.rows[0] == ["Çay Bardağı", "₺49,90", "Stokta", f"{site.base}/urun/1"]
    assert r.rows[1][0] == "Şeker Kasesi"  # fazla boşluk sadeleştirildi
    assert r.rows[1][3] == f"{site.base}/urun/2"  # göreli bağlantı sayfaya göre çözüldü
    assert r.rows[2][3] == "https://baska.test/u/3"
    assert r.status_code == 200 and r.engine == "static"


def test_links_and_text_modes(site):
    r = scrape_with_selector(f"{site.base}/urunler.html", "nav", "links", options=_opts())
    assert r.rows == [["1", f"{site.base}/hakkimizda", "Hakkımızda"], ["2", f"{site.base}/iletisim", "İletişim"]]
    r = scrape_with_selector(f"{site.base}/urunler.html", ".price", "text", options=_opts(max_rows=2))
    assert r.rows == [["1", "₺49,90"], ["2", "₺120,00"]]


def test_pagination_parses_every_page_and_stops_on_loop(site):
    url = f"{site.base}/liste-1.html"
    r = scrape_with_selector(url, "table.veri", "table", options=_opts(follow_pagination=True))
    assert r.columns == ["Şehir", "Nüfus"]
    assert [row[0] for row in r.rows] == ["İstanbul", "Ankara", "İzmir"]  # 2. sayfanın tablosu da alındı
    assert site.requested("/liste-1.html") == 1 and site.requested("/liste-2.html") == 1

    r = scrape_with_selector(url, "ul.haber", "links", options=_opts(follow_pagination=True))
    assert r.rows == [["1", f"{site.base}/detay/a", "Haber A"], ["2", f"{site.base}/detay/b", "Haber B"]]


def test_pagination_respects_delay_between_requests(site):
    r = scrape_with_selector(
        f"{site.base}/liste-1.html", "table.veri", "table", options=_opts(follow_pagination=True, delay_ms=300)
    )
    assert len(r.rows) == 3
    times = [t for p, t in site.log if p.startswith("/liste-")]
    assert times[1] - times[0] >= 0.29


def test_robots_disallow_blocks_before_fetching(site):
    site.routes["/robots.txt"] = (200, "User-agent: *\nDisallow: /urunler.html\n")
    with pytest.raises(ValueError, match="robots.txt"):
        scrape_with_selector(f"{site.base}/urunler.html", ".product", "text", options=_opts(robots_check=True))
    assert site.requested("/urunler.html") == 0
    # Tarayıcı modunda da (HTML zaten elde) çıkarma reddedilir
    with pytest.raises(ValueError, match="robots.txt"):
        scrape_from_html("<p>x</p>", f"{site.base}/urunler.html", "p", "text", options=_opts(robots_check=True))


def test_robots_partial_disallow_and_missing_file_allow(site):
    site.routes["/robots.txt"] = (200, "User-agent: *\nDisallow: /ozel/\n")
    r = scrape_with_selector(f"{site.base}/urunler.html", ".price", "text", options=_opts(robots_check=True))
    assert r.robots_warning is None and len(r.rows) == 3

    del site.routes["/robots.txt"]  # 404 -> robots.txt yok, serbest
    r = scrape_with_selector(f"{site.base}/urunler.html", ".price", "text", options=_opts(robots_check=True))
    assert r.robots_warning is None

    site.routes["/robots.txt"] = (500, "hata")  # okunamadı -> uyarı ama devam
    r = scrape_with_selector(f"{site.base}/urunler.html", ".price", "text", options=_opts(robots_check=True))
    assert r.robots_warning and "okunamadı" in r.robots_warning


def test_robots_crawl_delay_is_honoured(site):
    site.routes["/robots.txt"] = (200, "User-agent: *\nCrawl-delay: 1\n")
    t0 = time.monotonic()
    scrape_with_selector(
        f"{site.base}/liste-1.html", "table.veri", "table", options=_opts(robots_check=True, follow_pagination=True)
    )
    assert time.monotonic() - t0 >= 0.95


def test_robots_check_off_never_requests_robots(site):
    scrape_with_selector(f"{site.base}/urunler.html", ".price", "text", options=_opts(robots_check=False))
    assert site.requested("/robots.txt") == 0


def test_errors_are_turkish_value_errors(site):
    with pytest.raises(ValueError, match="Geçerli bir URL"):
        scrape_with_selector("ftp://x", "p", "text", options=_opts())
    with pytest.raises(ValueError, match="Bağlantı hatası"):
        scrape_with_selector("http://127.0.0.1:9/", "p", "text", options=_opts(timeout_s=2))
    with pytest.raises(ValueError, match="tablo bulunamadı"):
        scrape_with_selector(f"{site.base}/urunler.html", ".product", "table", options=_opts())
    with pytest.raises(ValueError, match="HTML boş"):
        scrape_from_html("  ", "http://x.test", "p", "text", options=_opts())


@pytest.fixture
def result(site):
    cols = {"name": "strong", "price": ".price", "stock": ".stock", "link": "a@href"}
    return scrape_with_selector(f"{site.base}/urunler.html", ".product", "table", options=_opts(columns=cols))


@pytest.mark.parametrize("encoding,read_enc", [("utf-8", "utf-8"), ("utf-8-bom", "utf-8-sig"), ("iso-8859-9", "iso-8859-9")])
def test_export_csv_roundtrip(tmp_path, result, encoding, read_enc):
    if encoding == "iso-8859-9":  # ₺ bu kodlamada yok; Türkçe harfler var
        result.rows = [[c.replace("₺", "TL ") for c in row] for row in result.rows]
    path = tmp_path / "out.csv"
    export_csv(str(path), result, encoding=encoding)
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf") == (encoding == "utf-8-bom")
    with open(path, newline="", encoding=read_enc) as f:
        rows = list(csv.reader(f))
    assert rows[0] == result.columns
    assert rows[1:] == result.rows


def test_export_csv_unencodable_char_raises(tmp_path, result):
    # UI bunu yakalayıp "Kaydedilemedi" gösterir; sessiz bozuk dosya yok
    with pytest.raises(UnicodeEncodeError):
        export_csv(str(tmp_path / "x.csv"), result, encoding="iso-8859-9")


def test_export_json_roundtrip(tmp_path, result):
    path = tmp_path / "out.json"
    export_json(str(path), result)
    data = json.loads(path.read_text(encoding="utf-8"))
    assert data["columns"] == result.columns and data["rows"] == result.rows
    assert data["meta"]["selector"] == ".product" and data["meta"]["mode"] == "table"
    assert "Çay" in path.read_text(encoding="utf-8")  # ensure_ascii=False


def test_export_xlsx_roundtrip_keeps_formula_text_as_text(tmp_path, result):
    from openpyxl import load_workbook

    path = tmp_path / "out.xlsx"
    export_xlsx(str(path), result)
    ws = load_workbook(path).active
    values = [[c if c is not None else "" for c in row] for row in ws.iter_rows(values_only=True)]
    assert values[0] == result.columns
    assert values[1:] == result.rows
    formula_cell = ws.cell(row=4, column=1)
    assert formula_cell.value.startswith("=HYPERLINK") and formula_cell.data_type == "s"
