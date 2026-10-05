from __future__ import annotations

import csv
import json
import re
import threading
import time
from dataclasses import dataclass, field
from html.parser import HTMLParser
from urllib.error import HTTPError
from urllib.parse import urljoin, urlparse
from urllib.request import Request, urlopen
from urllib.robotparser import RobotFileParser

from utils.config import resolve_user_agent


class _LinkParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.title = ""
        self._in_title = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "title":
            self._in_title = True

    def handle_endtag(self, tag: str) -> None:
        if tag == "title":
            self._in_title = False

    def handle_data(self, data: str) -> None:
        if self._in_title:
            self.title += data


@dataclass
class ScrapeOptions:
    max_rows: int = 500
    timeout_s: int = 30
    delay_ms: int = 0
    user_agent: str = "Mozilla/5.0 (compatible; LocalScraper/1.0)"
    strip_html: bool = True
    follow_pagination: bool = False
    robots_check: bool = True
    columns: dict[str, str] | None = None
    attr_mode: str = "text"
    accept_language: str = "tr-TR,tr;q=0.9,en;q=0.8"
    engine: str = "static"


@dataclass
class ScrapeResult:
    mode: str
    columns: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    title: str = ""
    robots_warning: str | None = None
    status_code: int = 200
    size_kb: float = 0.0
    elapsed_ms: int = 0
    url: str = ""
    selector: str = ""
    engine: str = "static"


def fetch_html(
    url: str,
    *,
    user_agent: str,
    timeout_s: int = 30,
    max_bytes: int = 1_500_000,
    accept_language: str = "",
) -> tuple[str, int, float]:
    parsed = urlparse(url.strip())
    if parsed.scheme not in ("http", "https"):
        raise ValueError("Geçerli bir URL girin (https://…)")

    headers = {"User-Agent": user_agent}
    if accept_language:
        headers["Accept-Language"] = accept_language
    req = Request(url, headers=headers)
    try:
        with urlopen(req, timeout=timeout_s) as resp:
            raw = resp.read(max_bytes)
            return raw.decode("utf-8", errors="replace"), getattr(resp, "status", 200), len(raw) / 1024
    except OSError as exc:  # URLError + zaman aşımı/bağlantı kopması (TimeoutError vb.)
        raise ValueError(f"Bağlantı hatası: {exc}") from exc


def _load_robots(url: str, *, user_agent: str, timeout_s: int = 8) -> tuple[RobotFileParser | None, str | None]:
    """(parser, uyarı). robots.txt yoksa (4xx) parser boş = her şey serbest (RFC 9309)."""
    parsed = urlparse(url.strip())
    if not parsed.netloc:
        return None, None
    robots_url = f"{parsed.scheme}://{parsed.netloc}/robots.txt"
    rp = RobotFileParser(robots_url)
    req = Request(robots_url, headers={"User-Agent": user_agent})
    try:
        with urlopen(req, timeout=timeout_s) as resp:
            lines = resp.read(512_000).decode("utf-8", errors="replace").splitlines()
    except HTTPError as exc:
        if 400 <= exc.code < 500:
            rp.parse([])
            return rp, None
        return None, "robots.txt okunamadı — yalnızca izinli sitelerden veri çıkarın."
    except OSError:
        return None, "robots.txt okunamadı — yalnızca izinli sitelerden veri çıkarın."
    rp.parse(lines)
    return rp, None


def _ensure_allowed(rp: RobotFileParser | None, url: str, user_agent: str) -> None:
    if rp is not None and not rp.can_fetch(user_agent, url):
        raise ValueError(
            f"robots.txt bu sayfaya erişimi engelliyor: {url}\n"
            "Site sahibinin iznini alın; Gelişmiş > 'robots.txt otomatik kontrol' yalnızca izinli sitelerde kapatılmalı."
        )


def check_robots_txt(url: str, *, user_agent: str) -> tuple[str | None, float]:
    """(uyarı, crawl-delay sn). Sayfa robots.txt ile engelliyse ValueError."""
    rp, warning = _load_robots(url, user_agent=user_agent)
    _ensure_allowed(rp, url, user_agent)
    delay = float(rp.crawl_delay(user_agent) or 0) if rp else 0.0
    return warning, delay


# Aynı sunucuya ardışık istekler arasında en az `delay_s` bekle (çalıştırmalar arası da geçerli).
_last_hit: dict[str, float] = {}
_throttle_lock = threading.Lock()


def _throttle(url: str, delay_s: float) -> None:
    host = urlparse(url).netloc.lower()
    with _throttle_lock:  # ponytail: global kilit; paralel çok-host kazıma gerekirse host başına kilit
        wait = _last_hit.get(host, -1e9) + delay_s - time.monotonic()
        if wait > 0:
            time.sleep(wait)
        _last_hit[host] = time.monotonic()


def _soup(html: str):
    try:
        from bs4 import BeautifulSoup
    except ImportError as exc:
        raise ValueError("beautifulsoup4 gerekli: pip install beautifulsoup4") from exc
    return BeautifulSoup(html, "html.parser")


def _page_title(html: str) -> str:
    parser = _LinkParser()
    parser.feed(html)
    return parser.title.strip() or "(yok)"


def _extract_value(el, spec: str, attr_mode: str, *, strip_html: bool) -> str:
    spec = (spec or "").strip()
    if not spec:
        return ""

    css, attr = spec, ""
    if "@" in spec:
        css, attr = spec.rsplit("@", 1)
        css, attr = css.strip(), attr.strip()

    target = el.select_one(css) if css else el
    if not target:
        return ""

    if attr:
        return str(target.get(attr, "") or "").strip()[:500]

    if attr_mode == "href":
        val = target.get("href", "") or target.get_text(strip=True)
    elif attr_mode == "src":
        val = target.get("src", "") or target.get_text(strip=True)
    elif attr_mode == "data-id":
        val = target.get("data-id", "") or target.get_text(strip=True)
    else:
        val = target.get_text(separator=" ", strip=not strip_html)
        if strip_html:
            val = re.sub(r"\s+", " ", val).strip()

    return str(val).strip()[:500]


def _parse_rows(
    html: str,
    url: str,
    selector: str,
    mode: str,
    opts: ScrapeOptions,
) -> tuple[list[str], list[list[str]]]:
    soup = _soup(html)
    sel = selector.strip()
    elements = soup.select(sel) if sel else [soup]
    mode = mode.lower()

    if mode == "links":
        cols = ["#", "Bağlantı", "Metin"]
        rows: list[list[str]] = []
        for el in elements:
            anchors = [el] if el.name == "a" and el.get("href") else el.find_all("a", href=True)
            for a in anchors:
                rows.append([a.get("href", ""), a.get_text(strip=True)])
        rows = rows[: opts.max_rows]
        return cols, [[str(i), urljoin(url, href), text] for i, (href, text) in enumerate(rows, 1)]

    if mode == "text":
        cols = ["#", "Metin"]
        out = []
        for i, el in enumerate(elements[: opts.max_rows], 1):
            t = el.get_text(separator=" ", strip=True)
            if t:
                out.append([str(i), t[:500]])
        return cols, out

    if mode == "table":
        if opts.columns and any(opts.columns.values()):
            headers, keys = [], []
            labels = {
                "name": "Ürün / başlık",
                "price": "Fiyat / değer",
                "stock": "Stok / meta",
                "link": "Bağlantı",
            }
            for key, label in labels.items():
                spec = (opts.columns or {}).get(key, "").strip()
                if spec or key == "name":
                    headers.append(label)
                    keys.append(key)
            rows = []
            for el in elements[: opts.max_rows]:
                row = [
                    _extract_value(el, (opts.columns or {}).get(k, ""), opts.attr_mode, strip_html=opts.strip_html)
                    for k in keys
                ]
                if any(cell for cell in row):
                    if "link" in keys:
                        idx = keys.index("link")
                        if row[idx]:
                            row[idx] = urljoin(url, row[idx])
                    rows.append(row)
            return headers, rows

        tables = []
        for el in elements:
            if el.name == "table":
                tables.append(el)
            else:
                tables.extend(el.find_all("table"))
        if not tables:
            raise ValueError("Seçicide tablo bulunamadı — sütun eşlemesi veya table seçici deneyin.")

        table = tables[0]
        headers = [th.get_text(strip=True) for th in table.find_all("th")]
        if not headers:
            first = table.find("tr")
            headers = [f"col{i+1}" for i in range(len(first.find_all(["td", "th"])))] if first else ["col1"]
        rows = []
        for tr in table.find_all("tr")[: opts.max_rows]:
            cells = [td.get_text(strip=True) for td in tr.find_all(["td", "th"])]
            if cells and cells != headers:
                while len(cells) < len(headers):
                    cells.append("")
                rows.append(cells[: len(headers)])
        return headers, rows

    raise ValueError(f"Bilinmeyen mod: {mode}")


def scrape_from_html(
    html: str,
    url: str,
    selector: str,
    mode: str,
    *,
    options: ScrapeOptions | None = None,
    settings: dict | None = None,
) -> ScrapeResult:
    opts = options or ScrapeOptions()
    if settings:
        opts.user_agent = resolve_user_agent(settings)
        opts.max_rows = int(settings.get("max_rows", opts.max_rows))
        opts.strip_html = bool(settings.get("strip_html", opts.strip_html))
        opts.robots_check = bool(settings.get("robots_check", opts.robots_check))
        opts.accept_language = settings.get("accept_language", opts.accept_language)
        opts.engine = settings.get("fetch_engine", "browser")

    t0 = time.perf_counter()
    if not html.strip():
        raise ValueError("Sayfa HTML boş — tarayıcıda sayfayı açıp tekrar deneyin.")

    warning = check_robots_txt(url, user_agent=opts.user_agent)[0] if opts.robots_check else None
    title = _page_title(html)
    cols, rows = _parse_rows(html, url, selector, mode, opts)
    elapsed_ms = int((time.perf_counter() - t0) * 1000)
    size_kb = len(html.encode("utf-8", errors="replace")) / 1024

    return ScrapeResult(
        mode=mode.lower(),
        columns=cols,
        rows=rows,
        title=title,
        robots_warning=warning,
        status_code=200,
        size_kb=size_kb,
        elapsed_ms=elapsed_ms,
        url=url,
        selector=selector.strip(),
        engine=opts.engine,
    )


def count_in_html(html: str, selector: str) -> int:
    return len(_soup(html).select(selector.strip()))


def _collect_pages(
    start_url: str,
    *,
    user_agent: str,
    timeout_s: int,
    follow_pagination: bool,
    accept_language: str = "",
    max_pages: int = 10,
    delay_s: float = 0.0,
    robots: RobotFileParser | None = None,
) -> tuple[list[tuple[str, str]], int, float]:
    """[(sayfa_url, html), ...], son durum kodu, toplam KB."""
    pages: list[tuple[str, str]] = []
    total_kb = 0.0
    status = 200
    url = start_url
    seen: set[str] = set()

    for _ in range(max_pages):
        if url in seen:
            break
        seen.add(url)
        _ensure_allowed(robots, url, user_agent)
        _throttle(url, delay_s)
        html, status, kb = fetch_html(
            url, user_agent=user_agent, timeout_s=timeout_s, accept_language=accept_language
        )
        pages.append((url, html))
        total_kb += kb
        if not follow_pagination:
            break
        nxt = _soup(html).select_one('a[rel="next"][href]')
        if not nxt:
            break
        url = urljoin(url, nxt["href"])

    return pages, status, total_kb


def scrape_with_selector(
    url: str,
    selector: str,
    mode: str,
    *,
    options: ScrapeOptions | None = None,
    settings: dict | None = None,
    html: str | None = None,
) -> ScrapeResult:
    # İki yol: html verilmişse (tarayıcı modu, Qt WebEngine'in render ettiği DOM) doğrudan
    # ayrıştırılır; yoksa urllib ile indirilir (statik mod, isteğe bağlı rel="next" takibi).
    opts = options or ScrapeOptions()
    if settings:
        opts.user_agent = resolve_user_agent(settings)
        opts.timeout_s = int(settings.get("timeout_s", opts.timeout_s))
        opts.delay_ms = int(settings.get("delay_ms", opts.delay_ms))
        opts.max_rows = int(settings.get("max_rows", opts.max_rows))
        opts.strip_html = bool(settings.get("strip_html", opts.strip_html))
        opts.follow_pagination = bool(settings.get("follow_pagination", opts.follow_pagination))
        opts.robots_check = bool(settings.get("robots_check", opts.robots_check))
        opts.accept_language = settings.get("accept_language", opts.accept_language)
        opts.engine = settings.get("fetch_engine", "static")

    if html is not None:
        opts.engine = settings.get("fetch_engine", "browser") if settings else "browser"
        return scrape_from_html(html, url, selector, mode, options=opts, settings=settings)

    t0 = time.perf_counter()
    robots, warning, crawl_delay = None, None, 0.0
    if opts.robots_check:
        robots, warning = _load_robots(url, user_agent=opts.user_agent)
        crawl_delay = float(robots.crawl_delay(opts.user_agent) or 0) if robots else 0.0
    pages, status, size_kb = _collect_pages(
        url,
        user_agent=opts.user_agent,
        timeout_s=opts.timeout_s,
        follow_pagination=opts.follow_pagination,
        accept_language=opts.accept_language,
        delay_s=max(opts.delay_ms / 1000, crawl_delay),
        robots=robots,
    )
    elapsed_ms = int((time.perf_counter() - t0) * 1000)

    # Her sayfa ayrı ayrıştırılır: göreli bağlantılar kendi sayfasına göre çözülür, tablo modu
    # yalnızca ilk sayfanın tablosunu almaz.
    cols: list[str] = []
    rows: list[list[str]] = []
    for page_url, page_html in pages:
        page_cols, page_rows = _parse_rows(page_html, page_url, selector, mode, opts)
        cols = cols or page_cols
        rows.extend(page_rows)
    rows = rows[: opts.max_rows]
    if cols and cols[0] == "#":
        rows = [[str(i), *r[1:]] for i, r in enumerate(rows, 1)]

    return ScrapeResult(
        mode=mode.lower(),
        columns=cols,
        rows=rows,
        title=_page_title(pages[0][1]),
        robots_warning=warning,
        status_code=status,
        size_kb=size_kb,
        elapsed_ms=elapsed_ms,
        url=url,
        selector=selector.strip(),
        engine="static",
    )


def count_selector_matches(url: str, selector: str, *, settings: dict, html: str | None = None) -> int:
    if html is not None:
        return count_in_html(html, selector)
    ua = resolve_user_agent(settings)
    page_html, _, _ = fetch_html(
        url,
        user_agent=ua,
        timeout_s=int(settings.get("timeout_s", 30)),
        accept_language=settings.get("accept_language", ""),
    )
    return count_in_html(page_html, selector)


def export_csv(path: str, result: ScrapeResult, *, encoding: str = "utf-8") -> None:
    enc = "utf-8-sig" if encoding == "utf-8-bom" else encoding
    with open(path, "w", newline="", encoding=enc) as f:
        writer = csv.writer(f)
        if result.columns:
            writer.writerow(result.columns)
        writer.writerows(result.rows)


def export_json(path: str, result: ScrapeResult) -> None:
    payload = {
        "meta": {
            "url": result.url,
            "selector": result.selector,
            "mode": result.mode,
            "title": result.title,
            "engine": result.engine,
        },
        "columns": result.columns,
        "rows": result.rows,
    }
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)


def export_xlsx(path: str, result: ScrapeResult) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sonuç"
    if result.columns:
        ws.append(result.columns)
        ws.freeze_panes = "A2"
    for row in result.rows:
        ws.append(row)
        for cell in ws[ws.max_row]:
            if cell.data_type == "f":  # "=..." metni formül olarak çalışmasın, düz metin kalsın
                cell.data_type = "s"
    wb.save(path)
